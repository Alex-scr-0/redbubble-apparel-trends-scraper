"""Rate-limited Redbubble apparel research scraper and Excel report builder."""

from __future__ import annotations

import json
import logging
import math
import os
import re
import time
from collections import defaultdict
from datetime import datetime, timezone
from typing import Any
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry


# Configuration
SEARCH_QUERIES = ["funny shirt", "gaming shirt", "cat shirt", "fitness shirt"]
MAX_PAGES_PER_QUERY = 2
MAX_PRODUCTS_TOTAL = 100
REQUEST_DELAY = 2.5
OUTPUT_FILENAME = os.path.join(
	os.path.expanduser('~'), 'Desktop', 'redbubble_clothing_research.xlsx'
)

REQUEST_TIMEOUT = 25
MAX_COLUMN_WIDTH = 60
BASE_URL = "https://www.redbubble.com"
USER_AGENT = "Mozilla/5.0 (compatible; ApparelMarketResearch/1.0; +https://www.redbubble.com/)"

APPAREL_CATEGORIES = {
	"t-shirts", "t-shirt", "tees", "hoodies", "sweatshirts", "sweaters",
	"crop-tops", "tank-tops", "long-sleeve-t-shirts", "long-sleeve-shirts",
	"dresses", "leggings", "skirts", "jackets", "kids-t-shirts",
	"baby-t-shirts", "onesies", "bodysuits", "shirts",
}
PRODUCT_PATH_PATTERN = re.compile(r"/i/([^/]+)/", re.IGNORECASE)
STOP_WORDS = {
	"a", "an", "and", "are", "for", "from", "in", "is", "it", "of",
	"on", "or", "the", "to", "with", "shirt", "shirts", "tee", "tees",
}

logging.basicConfig(
	level=logging.INFO,
	format="%(asctime)s %(levelname)s %(message)s",
)
LOGGER = logging.getLogger("redbubble_clothing_research")


def build_session() -> requests.Session:
	"""Create a reusable session with bounded retries for transient failures."""
	retry_policy = Retry(
		total=4,
		connect=4,
		read=3,
		status=4,
		backoff_factor=1.0,
		status_forcelist=(429, 500, 502, 503, 504),
		allowed_methods=frozenset({"GET"}),
		respect_retry_after_header=True,
		raise_on_status=False,
	)
	adapter = HTTPAdapter(max_retries=retry_policy, pool_connections=10, pool_maxsize=10)
	session = requests.Session()
	session.mount("https://", adapter)
	session.mount("http://", adapter)
	session.headers.update(
		{
			"User-Agent": USER_AGENT,
			"Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
			"Accept-Language": "en-US,en;q=0.8",
		}
	)
	return session


def fetch_soup(
	session: requests.Session,
	url: str,
	last_request_at: float | None,
) -> tuple[BeautifulSoup | None, float | None]:
	"""Fetch one page politely, returning None for recoverable page failures."""
	if last_request_at is not None:
		wait_seconds = REQUEST_DELAY - (time.monotonic() - last_request_at)
		if wait_seconds > 0:
			time.sleep(wait_seconds)
	try:
		response = session.get(url, timeout=REQUEST_TIMEOUT)
		requested_at = time.monotonic()
		response.raise_for_status()
		return BeautifulSoup(response.text, "html.parser"), requested_at
	except requests.RequestException as exc:
		LOGGER.warning("Request failed for %s: %s", url, exc)
		return None, time.monotonic()


def apparel_category(product_url: str) -> str | None:
	"""Return a recognized apparel category, rejecting every other item type."""
	match = PRODUCT_PATH_PATTERN.search(urlparse(product_url).path)
	if not match:
		return None
	category = match.group(1).lower()
	return category if category in APPAREL_CATEGORIES else None


def parse_count(value: Any) -> int | None:
	if value is None:
		return None
	digits = re.sub(r"[^0-9]", "", str(value))
	return int(digits) if digits else None


def parse_rating(value: Any) -> float | None:
	if value is None:
		return None
	match = re.search(r"\d+(?:\.\d+)?", str(value))
	if not match:
		return None
	rating = float(match.group())
	return rating if 0 <= rating <= 5 else None


def iter_json_objects(value: Any):
	if isinstance(value, dict):
		yield value
		for nested in value.values():
			yield from iter_json_objects(nested)
	elif isinstance(value, list):
		for nested in value:
			yield from iter_json_objects(nested)


def extract_public_signals(soup: BeautifulSoup) -> tuple[int, float | None]:
	"""Extract a public review count and rating from structured or visible data."""
	review_count = 0
	rating = None

	for script in soup.select('script[type="application/ld+json"]'):
		try:
			structured_data = json.loads(script.string or script.get_text())
		except (json.JSONDecodeError, TypeError):
			continue
		for item in iter_json_objects(structured_data):
			aggregate = item.get("aggregateRating")
			if not isinstance(aggregate, dict):
				continue
			rating = rating or parse_rating(aggregate.get("ratingValue"))
			review_count = max(
				review_count,
				parse_count(aggregate.get("reviewCount")) or 0,
				parse_count(aggregate.get("ratingCount")) or 0,
			)

	visible_text = soup.get_text(" ", strip=True)
	if rating is None:
		rating_match = re.search(
			r"(\d+(?:\.\d+)?)\s*(?:out of\s*5|/\s*5)", visible_text, re.IGNORECASE
		)
		if rating_match:
			rating = parse_rating(rating_match.group(1))
	review_matches = re.findall(
		r"([\d,]+)\s+(?:customer\s+)?(?:reviews|ratings)\b",
		visible_text,
		re.IGNORECASE,
	)
	if review_matches:
		review_count = max(review_count, *(parse_count(match) or 0 for match in review_matches))
	return review_count, rating


def card_for_anchor(anchor):
	current = anchor
	for _ in range(6):
		if current is None or getattr(current, "name", None) in {"body", "html"}:
			break
		text = current.get_text(" ", strip=True)
		if len(text) >= 20:
			return current
		current = current.parent
	return anchor.parent or anchor


def product_title(anchor, card, product_url: str) -> str:
	candidates = [
		anchor.get("title"),
		anchor.get("aria-label"),
		anchor.get_text(" ", strip=True),
	]
	image = card.find("img") if hasattr(card, "find") else None
	if image:
		candidates.append(image.get("alt"))
	for candidate in candidates:
		cleaned = re.sub(r"\s+", " ", candidate or "").strip()
		if cleaned and len(cleaned) > 2 and not cleaned.lower().startswith("shop "):
			return cleaned[:300]
	slug = urlparse(product_url).path.rstrip("/").split("/")[-1]
	return re.sub(r"[-_]+", " ", slug)[:300] or "Untitled apparel product"


def parse_search_products(soup: BeautifulSoup) -> list[dict[str, Any]]:
	"""Read unique, explicitly apparel-categorized product links from a result page."""
	products: list[dict[str, Any]] = []
	seen_urls: set[str] = set()
	for anchor in soup.select('a[href*="/i/"]'):
		product_url = urljoin(BASE_URL, anchor.get("href", "")).split("?", 1)[0]
		if urlparse(product_url).netloc.lower() not in {"redbubble.com", "www.redbubble.com"}:
			continue
		category = apparel_category(product_url)
		if category is None or product_url in seen_urls:
			continue
		seen_urls.add(product_url)
		card = card_for_anchor(anchor)
		card_text = card.get_text(" ", strip=True)
		artist_match = re.search(r"(?:by|designed by)\s+([^|·]+)", card_text, re.IGNORECASE)
		products.append(
			{
				"product_url": product_url,
				"category": category,
				"title": product_title(anchor, card, product_url),
				"artist": artist_match.group(1).strip()[:150] if artist_match else "",
				"card_review_count": 0,
				"card_rating": None,
			}
		)
	return products


def popularity_score(review_count: int, rating: float | None, rank: int, max_rank: int) -> float:
	"""Score public signals on a 0-100 scale; missing ratings contribute zero.

	Formula: 40 * min(log(1 + reviews) / log(1001), 1)
	+ 30 * (rating / 5, or 0 when unavailable)
	+ 30 * (1 - (rank - 1) / (max_rank - 1)); the rank term is 30 when max_rank is 1.
	Review volume saturates at 1,000 reviews. Search position is normalized within
	the collected results for the product's primary query. The score is a research
	proxy, not Redbubble's internal ranking or sales data.
	"""
	review_points = 40 * min(math.log1p(max(review_count, 0)) / math.log1p(1000), 1)
	rating_points = 30 * ((rating or 0) / 5)
	rank_points = 30 if max_rank <= 1 else 30 * max(0, 1 - (rank - 1) / (max_rank - 1))
	return round(min(100, review_points + rating_points + rank_points), 2)


def scrape_products() -> list[dict[str, Any]]:
	session = build_session()
	products_by_url: dict[str, dict[str, Any]] = {}
	last_request_at = None

	try:
		for query in SEARCH_QUERIES:
			if len(products_by_url) >= MAX_PRODUCTS_TOTAL:
				break
			query_rank_offset = 0
			for page_number in range(1, MAX_PAGES_PER_QUERY + 1):
				if len(products_by_url) >= MAX_PRODUCTS_TOTAL:
					break
				search_url = f"{BASE_URL}/shop/"
				params = {"query": query, "page": page_number}
				try:
					if last_request_at is not None:
						wait_seconds = REQUEST_DELAY - (time.monotonic() - last_request_at)
						if wait_seconds > 0:
							time.sleep(wait_seconds)
					response = session.get(
						search_url, params=params, timeout=REQUEST_TIMEOUT
					)
					last_request_at = time.monotonic()
					response.raise_for_status()
					soup = BeautifulSoup(response.text, "html.parser")
				except requests.RequestException as exc:
					LOGGER.warning("Search failed for %r page %d: %s", query, page_number, exc)
					last_request_at = time.monotonic()
					continue

				page_products = parse_search_products(soup)
				LOGGER.info(
					"Query %r page %d: found %d apparel products",
					query,
					page_number,
					len(page_products),
				)
				for page_rank, candidate in enumerate(page_products, start=1):
					if len(products_by_url) >= MAX_PRODUCTS_TOTAL:
						break
					existing = products_by_url.get(candidate["product_url"])
					if existing:
						if query not in existing["matched_queries"]:
							existing["matched_queries"].append(query)
						continue

					detail_soup, last_request_at = fetch_soup(
						session, candidate["product_url"], last_request_at
					)
					review_count = candidate["card_review_count"]
					rating = candidate["card_rating"]
					if detail_soup is not None:
						detail_reviews, detail_rating = extract_public_signals(detail_soup)
						review_count = max(review_count, detail_reviews)
						rating = rating if rating is not None else detail_rating

					products_by_url[candidate["product_url"]] = {
						**candidate,
						"query": query,
						"matched_queries": [query],
						"page_number": page_number,
						"search_rank": query_rank_offset + page_rank,
						"review_count": review_count,
						"rating": rating,
						"scraped_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
					}
				query_rank_offset += len(page_products)
	finally:
		session.close()

	max_rank_by_query: dict[str, int] = defaultdict(int)
	for product in products_by_url.values():
		max_rank_by_query[product["query"]] = max(
			max_rank_by_query[product["query"]], product["search_rank"]
		)
	for product in products_by_url.values():
		product["popularity_score"] = popularity_score(
			product["review_count"],
			product["rating"],
			product["search_rank"],
			max_rank_by_query[product["query"]],
		)
	return list(products_by_url.values())


def average(values: list[float]) -> float | None:
	return round(sum(values) / len(values), 2) if values else None


def build_workbook(products: list[dict[str, Any]]) -> None:
	workbook = Workbook()
	products_sheet = workbook.active
	products_sheet.title = "Products"
	sheet_headers: dict[str, list[str]] = {
		"Products": [
			"Search Query", "Matched Queries", "Search Rank", "Page", "Product Title",
			"Artist", "Apparel Category", "Reviews", "Public Rating", "Popularity Signal Score",
			"Product URL", "Scraped UTC",
		],
		"Top Products": [
			"Rank", "Product Title", "Search Query", "Apparel Category", "Reviews",
			"Public Rating", "Popularity Signal Score", "Product URL",
		],
		"Niches": [
			"Niche Query", "Products", "Average Popularity Score", "Average Rating", "Review Total",
		],
		"Keywords": ["Keyword", "Product Mentions", "Average Popularity Score"],
		"Trends": [
			"Search Query", "Page", "Products", "Average Popularity Score",
			"Average Rating", "Review Total",
		],
		"Statistics": ["Metric", "Value"],
	}
	sheets = {"Products": products_sheet}
	for sheet_name in list(sheet_headers)[1:]:
		sheets[sheet_name] = workbook.create_sheet(sheet_name)

	product_rows = [
		[
			product["query"],
			", ".join(product["matched_queries"]),
			product["search_rank"],
			product["page_number"],
			product["title"],
			product["artist"],
			product["category"],
			product["review_count"],
			product["rating"],
			product["popularity_score"],
			product["product_url"],
			product["scraped_at"],
		]
		for product in products
	]
	top_products = sorted(
		products,
		key=lambda product: (product["popularity_score"], product["review_count"]),
		reverse=True,
	)[:25]
	top_rows = [
		[
			rank,
			product["title"],
			product["query"],
			product["category"],
			product["review_count"],
			product["rating"],
			product["popularity_score"],
			product["product_url"],
		]
		for rank, product in enumerate(top_products, start=1)
	]

	by_query: dict[str, list[dict[str, Any]]] = defaultdict(list)
	by_query_page: dict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
	keyword_products: dict[str, list[dict[str, Any]]] = defaultdict(list)
	for product in products:
		by_query[product["query"]].append(product)
		by_query_page[(product["query"], product["page_number"])].append(product)
		for word in set(re.findall(r"[a-z0-9]+", product["title"].lower())):
			if len(word) > 2 and word not in STOP_WORDS:
				keyword_products[word].append(product)

	niche_rows = []
	for query in SEARCH_QUERIES:
		group = by_query.get(query, [])
		niche_rows.append(
			[
				query,
				len(group),
				average([item["popularity_score"] for item in group]),
				average([item["rating"] for item in group if item["rating"] is not None]),
				sum(item["review_count"] for item in group),
			]
		)

	keyword_rows = [
		[word, len(items), average([item["popularity_score"] for item in items])]
		for word, items in sorted(
			keyword_products.items(),
			key=lambda pair: (len(pair[1]), pair[0]),
			reverse=True,
		)
	]
	trend_rows = []
	for (query, page_number), group in sorted(by_query_page.items()):
		trend_rows.append(
			[
				query,
				page_number,
				len(group),
				average([item["popularity_score"] for item in group]),
				average([item["rating"] for item in group if item["rating"] is not None]),
				sum(item["review_count"] for item in group),
			]
		)

	rated_products = [item["rating"] for item in products if item["rating"] is not None]
	statistics_rows = [
		["Products collected", len(products)],
		["Queries configured", len(SEARCH_QUERIES)],
		["Maximum pages per query", MAX_PAGES_PER_QUERY],
		["Average popularity score", average([item["popularity_score"] for item in products])],
		["Products with a public rating", len(rated_products)],
		["Average public rating", average(rated_products)],
		["Total publicly visible reviews", sum(item["review_count"] for item in products)],
		["Generated UTC", datetime.now(timezone.utc).isoformat(timespec="seconds")],
		["Scoring note", "Research proxy from public reviews, public rating, and collected search rank."],
	]

	data_by_sheet = {
		"Products": product_rows,
		"Top Products": top_rows,
		"Niches": niche_rows,
		"Keywords": keyword_rows,
		"Trends": trend_rows,
		"Statistics": statistics_rows,
	}
	for sheet_name, sheet in sheets.items():
		headers = sheet_headers[sheet_name]
		sheet.append(headers)
		for row in data_by_sheet[sheet_name]:
			sheet.append(row)
		sheet.freeze_panes = "A2"
		sheet.sheet_view.showGridLines = True
		sheet.auto_filter.ref = sheet.dimensions
		for cell in sheet[1]:
			cell.fill = PatternFill(fill_type="solid", fgColor="1B365D")
			cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
			cell.alignment = Alignment(vertical="center", wrap_text=True)
		sheet.row_dimensions[1].height = 30
		for column_cells in sheet.columns:
			column_letter = get_column_letter(column_cells[0].column)
			max_length = max(
				(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells),
				default=0,
			)
			sheet.column_dimensions[column_letter].width = min(
				max(max_length + 2, 12), MAX_COLUMN_WIDTH
			)
		for row in sheet.iter_rows(min_row=2):
			for cell in row:
				cell.alignment = Alignment(vertical="top", wrap_text=True)

	score_column = sheet_headers["Products"].index("Popularity Signal Score") + 1
	score_header = products_sheet.cell(row=1, column=score_column)
	score_header.comment = Comment(
		"Redbubble Popularity Signal Score (0-100), calculated from observable public signals. "
		"Formula: 40 * min(log(1 + review_count) / log(1001), 1) "
		"+ 30 * (public_rating / 5; zero when unavailable) "
		"+ 30 * (1 - (search_rank - 1) / (max_rank - 1)); rank points are 30 if max_rank is 1. "
		"Review points saturate at 1,000 reviews. Search rank is normalized over collected results "
		"for the primary query. This is a research proxy, not Redbubble's internal score or sales data.",
		"OpenAI",
	)
	workbook.save(OUTPUT_FILENAME)
	LOGGER.info("Saved %d apparel products to %s", len(products), OUTPUT_FILENAME)


def main() -> None:
	try:
		products = scrape_products()
		build_workbook(products)
	except (OSError, ValueError, requests.RequestException) as exc:
		LOGGER.exception("Scraper could not complete the report: %s", exc)
		raise
	except Exception:
		LOGGER.exception("Unexpected scraper failure")
		raise


if __name__ == "__main__":
	main()
